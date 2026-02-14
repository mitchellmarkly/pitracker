import json
import sys
from pathlib import Path
from openpyxl import load_workbook

def s(x):
    return "" if x is None else str(x).strip()

def n(x):
    if x is None or x == "":
        return None
    if isinstance(x, (int, float)):
        return float(x)
    try:
        return float(str(x).replace(",", "").strip())
    except Exception:
        return None

def convert(xlsx_path: Path, sheet_name: str, out_path: Path):
    wb = load_workbook(xlsx_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise SystemExit(f"Sheet not found: {sheet_name}. Available: {wb.sheetnames}")

    ws = wb[sheet_name]

    def cell(r, c):
        return ws.cell(row=r, column=c).value

    # 1-indexed rows/cols (openpyxl)
    # row 1: resources starting col 5
    resources = []
    col = 5
    while True:
        v = s(cell(1, col))
        if not v:
            break
        resources.append((col, v))
        col += 1

    if not resources:
        raise SystemExit("No resource headers found on row 1 starting at col 5.")

    region = sheet_name.split("-", 1)[1].strip() if "-" in sheet_name else sheet_name

    scans = []
    cur_const = ""
    cur_system = ""

    r = 3
    while True:
        const_cell = s(cell(r, 1))
        sys_cell = s(cell(r, 2))
        planet = s(cell(r, 3))
        ptype = s(cell(r, 4))

        if not const_cell and not sys_cell and not planet and not ptype:
            break

        if const_cell:
            cur_const = const_cell
        if sys_cell:
            cur_system = sys_cell
        if planet:
            for col, res in resources:
                val = n(cell(r, col))
                if val is None:
                    continue
                scans.append({
                    "Region": region,
                    "Constellation": cur_const,
                    "System": cur_system,
                    "Planet": planet,
                    "PlanetType": ptype,
                    "Resource": res,
                    "Value": val,
                })

        r += 1

    payload = {
        "schemaVersion": 1,
        "kind": "pi-heatmap",
        "region": region,
        "scans": scans
    }

    out_path.write_text(json.dumps(payload, indent=2))
    print(f"Wrote {out_path} ({len(scans)} scan rows)")

if __name__ == "__main__":
    if len(sys.argv) < 4:
        raise SystemExit("Usage: python convert_heatmap_from_xlsx.py <xlsx_path> <sheet_name> <out_json_path>")
    convert(Path(sys.argv[1]), sys.argv[2], Path(sys.argv[3]))
